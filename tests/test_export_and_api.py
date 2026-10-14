import io
import re

from docx import Document

from app.data_loader import load_template
from app.models import LessonRequest
from app.render.markdown_blocks import markdown_to_html


def _generate(client, code="VC2M4M02") -> str:
    response = client.post(
        "/generate",
        data={"curriculum_code": code, "student_level": "core", "duration_minutes": 60},
    )
    assert response.status_code == 200
    return re.search(r"/export/([a-f0-9]+)\.docx", response.text).group(1)


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    header = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    return header + "\n" + "\n".join(p.text for p in doc.paragraphs)


def test_index_page_lists_curriculum(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "VC2M5N05" in response.text
    assert "mock" in response.text


def test_generate_page_shows_sections_and_check_summary(client):
    response = client.post("/generate", data={"curriculum_code": "VC2M5ST01"})
    assert response.status_code == 200
    for section in load_template("default-tutoring").sections:
        assert section.title in response.text
    assert "道题答案已自动复算" in response.text
    assert "data:image/png;base64," in response.text


def test_teacher_export_has_branding_and_answers(client):
    plan_id = _generate(client)
    response = client.get(f"/export/{plan_id}.docx?variant=teacher")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    text = _docx_text(response.content)
    assert "Box Hill Bright Minds Tutoring" in text
    assert "Teacher lesson plan" in text
    assert "Answer:" in text and "Working:" in text
    assert "Note for Parents" in text


def test_student_export_hides_answers_and_teacher_sections(client):
    plan_id = _generate(client)
    text = _docx_text(client.get(f"/export/{plan_id}.docx?variant=student").content)
    assert "Student worksheet" in text
    assert "Tiered Practice" in text
    assert "Note for Parents" not in text
    assert "Teaching the New Idea" not in text
    practice_part = text.split("Tiered Practice", 1)[1]
    assert "Working:" not in practice_part
    assert "Answer: ____" in practice_part


def test_regenerate_endpoint_returns_updated_section(client):
    plan_id = _generate(client)
    response = client.post(
        f"/plans/{plan_id}/sections/warm_up/regenerate", data={"feedback": "Shorter please"}
    )
    assert response.status_code == 200
    assert 'id="section-warm_up"' in response.text
    assert "Shorter please" in response.text
    assert "已按反馈修改 1 次" in response.text


def test_unknown_ids_return_404(client):
    assert client.post("/generate", data={"curriculum_code": "NOPE"}).status_code == 404
    bad_org = {"curriculum_code": "VC2M5N05", "org_id": "../../etc"}
    assert client.post("/generate", data=bad_org).status_code == 404
    assert client.get("/export/doesnotexist.docx").status_code == 404
    plan_id = _generate(client)
    response = client.post(f"/plans/{plan_id}/sections/nope/regenerate", data={"feedback": "x"})
    assert response.status_code == 404


def test_json_api_round_trip(client):
    response = client.post(
        "/api/plans", json=LessonRequest(curriculum_code="VC2M6N07").model_dump(mode="json")
    )
    assert response.status_code == 200
    plan = response.json()
    assert client.get(f"/api/plans/{plan['id']}").json()["id"] == plan["id"]


def test_markdown_preview_escapes_html():
    html = markdown_to_html("## Title\n- **bold** <script>alert(1)</script>\n\nText _note_")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "<strong>bold</strong>" in html and "<em>note</em>" in html
    assert html.startswith("<h3>Title</h3><ul><li>")
