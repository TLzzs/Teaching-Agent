import io
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.data_loader import DATA_DIR
from app.models import (
    STUDENT_LEVEL_LABELS,
    CheckStatus,
    CurriculumItem,
    LessonPlan,
    OrgConfig,
    OutputType,
    Question,
    SectionResult,
    StudentLevel,
    WorkflowTemplate,
)
from app.render.diagrams import render_diagram, render_logo
from app.render.markdown_blocks import Run, parse_markdown

Variant = Literal["teacher", "student"]

TIER_TITLES = {
    StudentLevel.FOUNDATION: "Foundation",
    StudentLevel.CORE: "Core",
    StudentLevel.EXTENSION: "Extension",
}
CHECK_LABELS = {
    CheckStatus.VERIFIED: "Answer checked automatically",
    CheckStatus.MISMATCH: "CHECK FAILED - teacher must review",
    CheckStatus.INVALID: "Could not check automatically - teacher must review",
    CheckStatus.NOT_CHECKABLE: "Not auto-checkable - teacher to review",
}
REVIEW_NOTICE = "AI-generated draft. A teacher must review all content before use."


def _color(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _logo_bytes(org: OrgConfig) -> bytes:
    if org.logo_path:
        path = (DATA_DIR / "orgs" / org.logo_path).resolve()
        if path.is_relative_to(DATA_DIR) and path.is_file():
            return Path(path).read_bytes()
    return render_logo(org.name, org.primary_color)


def _add_runs(paragraph, runs: list[Run]) -> None:
    for run in runs:
        r = paragraph.add_run(run.text)
        r.bold = run.bold or None
        r.italic = run.italic or None


def _add_markdown(doc, text: str) -> None:
    for block in parse_markdown(text):
        if block.kind == "heading":
            paragraph = doc.add_paragraph(style="Heading 3")
        elif block.kind == "bullet":
            paragraph = doc.add_paragraph(style="List Bullet")
        elif block.kind == "numbered":
            paragraph = doc.add_paragraph(style="List Number")
        else:
            paragraph = doc.add_paragraph()
        _add_runs(paragraph, block.runs)


def _setup_document(org: OrgConfig):
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    brand = _color(org.primary_color)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = brand

    header = doc.sections[0].header.paragraphs[0]
    header.add_run().add_picture(io.BytesIO(_logo_bytes(org)), height=Cm(1.3))
    name_run = header.add_run(f"   {org.name}")
    name_run.bold = True
    name_run.font.size = Pt(13)
    name_run.font.color.rgb = brand
    if org.tagline:
        tagline = header.add_run(f"\n{org.tagline}")
        tagline.font.size = Pt(9)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(" | ".join(filter(None, [org.footer_text, REVIEW_NOTICE])))
    footer_run.font.size = Pt(8)
    return doc


def _add_info_table(doc, plan: LessonPlan, item: CurriculumItem) -> None:
    rows = [
        ("Curriculum", f"Victorian Curriculum 2.0 Mathematics - {item.strand}"),
        ("Content description", f"{item.code} (Level {item.level})"),
        ("Focus", item.summary),
        ("Student level", STUDENT_LEVEL_LABELS[plan.request.student_level]),
        ("Lesson length", f"{plan.request.duration_minutes} minutes"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = value
    doc.add_paragraph()


def _add_question(
    doc, number: int, question: Question, show_answers: bool, show_checks: bool
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{number}. ").bold = True
    paragraph.add_run(question.prompt)
    if question.diagram is not None:
        doc.add_paragraph().add_run().add_picture(
            io.BytesIO(render_diagram(question.diagram)), width=Cm(9)
        )
    if not show_answers:
        doc.add_paragraph("Answer: " + "_" * 40)
        return
    answer = doc.add_paragraph()
    answer.paragraph_format.left_indent = Cm(0.6)
    answer.add_run("Answer: ").bold = True
    answer.add_run(question.answer)
    if question.working:
        working = doc.add_paragraph()
        working.paragraph_format.left_indent = Cm(0.6)
        working.add_run("Working: ").italic = True
        working.add_run(question.working)
    if show_checks and question.check_result is not None:
        status = doc.add_paragraph()
        status.paragraph_format.left_indent = Cm(0.6)
        label = status.add_run(CHECK_LABELS[question.check_result.status])
        label.font.size = Pt(8)
        verified = question.check_result.status == CheckStatus.VERIFIED
        label.font.color.rgb = (
            RGBColor(0x15, 0x80, 0x3D) if verified else RGBColor(0xB9, 0x1C, 0x1C)
        )


def _add_questions(doc, section: SectionResult, show_answers: bool, show_checks: bool) -> None:
    questions = section.questions or []
    if any(q.tier for q in questions):
        number = 1
        for tier in StudentLevel:
            group = [q for q in questions if q.tier == tier]
            if not group:
                continue
            doc.add_heading(TIER_TITLES[tier], level=3)
            for question in group:
                _add_question(doc, number, question, show_answers, show_checks)
                number += 1
        untiered = [q for q in questions if q.tier is None]
    else:
        number, untiered = 1, questions
    for offset, question in enumerate(untiered):
        _add_question(doc, number + offset, question, show_answers, show_checks)


def build_docx(
    plan: LessonPlan,
    org: OrgConfig,
    item: CurriculumItem,
    template: WorkflowTemplate,
    variant: Variant = "teacher",
) -> bytes:
    """Teacher copy has every section with answers; student copy has worksheets only."""
    doc = _setup_document(org)
    title = f"Level {item.level} Mathematics: {item.topic}"
    doc.add_heading(title, level=0)
    subtitle = "Teacher lesson plan" if variant == "teacher" else "Student worksheet"
    doc.add_paragraph(subtitle).runs[0].italic = True

    if variant == "teacher":
        _add_info_table(doc, plan, item)
    else:
        doc.add_paragraph("Name: " + "_" * 30 + "    Date: " + "_" * 15)

    for section in plan.sections:
        definition = template.section(section.section_id)
        if variant == "student" and not definition.student_facing:
            continue
        doc.add_heading(section.title, level=1)
        if section.output_type == OutputType.MARKDOWN:
            _add_markdown(doc, section.markdown or "")
        else:
            teacher = variant == "teacher"
            _add_questions(
                doc,
                section,
                show_answers=teacher or definition.answers_in_student_copy,
                show_checks=teacher,
            )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
